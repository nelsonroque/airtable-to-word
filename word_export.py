"""Turn an Airtable schema and record into an editable Word form."""

from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO
from typing import Any

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


def format_value(value: Any, field_type: str = "") -> str:
    if value is None or value == "":
        return ""
    if isinstance(value, bool):
        return "☒ Yes" if value else "☐ No"
    if isinstance(value, list):
        if not value:
            return ""
        return "\n".join(format_value(item, field_type) for item in value)
    if isinstance(value, dict):
        if field_type == "multipleAttachments":
            filename = value.get("filename", "Attachment")
            return f"{filename}: {value.get('url', '')}".rstrip(": ")
        return str(
            value.get("name")
            or value.get("email")
            or value.get("url")
            or value.get("id")
            or value
        )
    return str(value)


def _shade_cell(cell: Any, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def _set_cell_margins(cell: Any, top: int = 100, start: int = 120, bottom: int = 100, end: int = 120) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")
        margins.append(element)


def _autofit_table_to_window(table: Any) -> None:
    """Match Word's AutoFit Window behavior (100% of usable page width)."""
    table.autofit = True
    properties = table._tbl.tblPr
    width = properties.first_child_found_in("w:tblW")
    if width is None:
        width = OxmlElement("w:tblW")
        properties.insert(0, width)
    width.set(qn("w:type"), "pct")
    width.set(qn("w:w"), "5000")


def build_document(table: dict[str, Any], record: dict[str, Any]) -> BytesIO:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run(f"{table.get('name', 'Airtable Record')} Form")
    title_run.bold = True
    title_run.font.size = Pt(18)
    title_run.font.color.rgb = RGBColor(31, 78, 121)

    metadata = document.add_paragraph()
    metadata.alignment = WD_ALIGN_PARAGRAPH.CENTER
    metadata_run = metadata.add_run(
        f"Record ID: {record.get('id', '')}  •  Generated: "
        f"{datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}"
    )
    metadata_run.italic = True
    metadata_run.font.size = Pt(8)
    metadata_run.font.color.rgb = RGBColor(95, 95, 95)

    fields = table.get("fields", [])
    values = record.get("fields", {})
    form = document.add_table(rows=1, cols=2)
    form.alignment = WD_TABLE_ALIGNMENT.CENTER
    form.style = "Table Grid"
    _autofit_table_to_window(form)

    header = form.rows[0].cells
    header[0].text = "Field"
    header[1].text = "Value (editable)"
    for cell in header:
        _shade_cell(cell, "1F4E79")
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)

    for field in fields:
        row = form.add_row().cells
        row[0].text = str(field.get("name", ""))
        row[1].text = format_value(values.get(field.get("name")), field.get("type", ""))
        row[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        row[1].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        _shade_cell(row[0], "D9EAF7")
        row[0].paragraphs[0].runs[0].bold = True
        for cell in row:
            _set_cell_margins(cell)
        if not row[1].text:
            row[1].paragraphs[0].add_run(" ")

        description = field.get("description")
        if description:
            hint = row[0].add_paragraph(str(description))
            hint.runs[0].italic = True
            hint.runs[0].font.size = Pt(7)
            hint.runs[0].font.color.rgb = RGBColor(95, 95, 95)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Generated from Airtable • Values in this document do not update Airtable automatically").font.size = Pt(7)

    output = BytesIO()
    document.save(output)
    output.seek(0)
    return output


def safe_filename(value: str) -> str:
    cleaned = "".join(char if char.isalnum() or char in "-_" else "_" for char in value)
    return cleaned.strip("_")[:80] or "airtable_record"
