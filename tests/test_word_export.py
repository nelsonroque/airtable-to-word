from docx import Document

from word_export import build_document, format_value, safe_filename


def test_format_value_handles_common_airtable_values():
    assert format_value(True) == "☒ Yes"
    assert format_value(["One", "Two"]) == "One\nTwo"
    assert format_value({"name": "Ada"}) == "Ada"
    assert format_value({"filename": "a.pdf", "url": "https://example.test/a.pdf"}, "multipleAttachments") == "a.pdf: https://example.test/a.pdf"


def test_document_includes_empty_schema_fields():
    table = {
        "name": "People",
        "fields": [
            {"id": "fld1", "name": "Name", "type": "singleLineText"},
            {"id": "fld2", "name": "Notes", "type": "multilineText"},
        ],
    }
    record = {"id": "rec123", "fields": {"Name": "Ada Lovelace"}}
    output = build_document(table, record)
    document = Document(output)
    cells = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    assert "Name" in cells
    assert "Ada Lovelace" in cells
    assert "Notes" in cells
    table_xml = document.tables[0]._tbl.xml
    assert 'w:type="pct"' in table_xml
    assert 'w:w="5000"' in table_xml
    assert 'w:type="autofit"' in table_xml


def test_safe_filename():
    assert safe_filename("Client / Example") == "Client___Example"
